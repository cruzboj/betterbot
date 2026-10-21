"""Build a slide deck with one slide per text: the rolling-mean plot (same as per_text/) above
the per-word plot (same as per_text_raw/), drawn as one figure with a shared x-axis so the two
line up. Figures go to <results>/slides_img/ (the existing plot folders are not touched).
Output: LCA_per_text_plots.pptx (opens in Google Slides / PowerPoint).

    python make_slides.py [results/k5_clean]
"""
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).parent
RES = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "results" / "k5_clean"
OUT = ROOT / "LCA_per_text_plots.pptx"

NAVY, INK, MUTED, WHITE = "1F2A44", "1F1F1F", "5F6670", "FFFFFF"
KIND_COLOR = {"factual": "2A78D6", "creative": "EB6834"}  # same blue / orange as the plots
FONT_HEAD, FONT_BODY = "Cambria", "Calibri"

W, H = 13.333, 7.5
PLOT_W = 12.0
MODELS = ["gpt2-medium", "EleutherAI/pythia-410m", "Qwen/Qwen2.5-1.5B"]
MODEL_COLORS = ["#2a78d6", "#eb6834", "#1baf7a"]  # same as the notebook's plots
AVG_COLOR, SMOOTH, LABEL = "#1f1f1f", 10, "LCA height (max)"


def load_table(name):
    """rows: content word (word_idx), cols: model -> lca_height; plus the rows' POS."""
    dfs = {m: pd.read_csv(RES / f"{m.replace('/', '_')}__{name}.csv") for m in MODELS}
    table = pd.DataFrame({m: d.set_index("word_idx")["lca_height"] for m, d in dfs.items()}).sort_index()
    pos = pd.concat(d.set_index("word_idx")["pos"] for d in dfs.values()).groupby(level=0).first()
    return table, pos


def figure(name, table, out):
    """Top: rolling mean (as per_text/). Bottom: every word (as per_text_raw/). Shared x-axis."""
    roll = lambda s: s.rolling(SMOOTH, center=True, min_periods=1).mean()
    fig, (top, bot) = plt.subplots(2, 1, figsize=(14, 6.9), sharex=True)
    for (m, s), c in zip(table.items(), MODEL_COLORS):
        top.plot(table.index, s, color=c, lw=0.7, alpha=0.18)
        top.plot(table.index, roll(s), color=c, lw=1.6, alpha=0.9, label=f"{m}  (mean {s.mean():.2f})")
        s = s.dropna()
        bot.plot(s.index, s, color=c, lw=0.9, alpha=0.75, marker="o", ms=2.5)
    avg = table.mean(axis=1)
    top.plot(table.index, roll(avg), color=AVG_COLOR, lw=3, label="average of models")
    bot.plot(table.index, avg, color=AVG_COLOR, lw=1.8)
    top.set_title(f"Rolling mean over {SMOOTH} words (thin lines: raw values)", loc="left", fontsize=11)
    bot.set_title("Every measured word (no smoothing)", loc="left", fontsize=11)
    for ax in (top, bot):
        ax.set_ylim(0, 17)
        ax.set_ylabel(LABEL)
        ax.grid(alpha=0.25, lw=0.6)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
    bot.set_xlabel("content word # along the text")
    top.legend(loc="upper left", fontsize=9, frameon=False, ncol=4)
    fig.tight_layout()
    fig.savefig(out, dpi=150)
    plt.close(fig)


def pretty(name):
    if name.startswith("factual_wiki_"):
        return "Wikipedia — " + name.removeprefix("factual_wiki_").replace("_", " ").capitalize()
    if name == "creative_harry_potter":
        return "Harry Potter — Book 1 opening"
    if name.startswith("creative_hp"):
        return f"Harry Potter — Book {name.removeprefix('creative_hp')} (mid-book excerpt)"
    if name == "creative_jokes_for_all_occasions":
        return "Jokes For All Occasions — excerpt 1"
    if name == "creative_jokes_2":
        return "Jokes For All Occasions — excerpt 2"
    return name


def order_key(name):
    kind = 0 if name.startswith("factual") else 1
    rank = {"creative_harry_potter": 0, "creative_jokes_for_all_occasions": 20, "creative_jokes_2": 21}
    return kind, rank.get(name, 10), name


def text(slide, x, y, w, h, s, size, color, bold=False, font=FONT_BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
    r.font.color.rgb = RGBColor.from_string(color)
    return tb


def pill(slide, x, y, w, h, label, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = RGBColor.from_string(color)
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size, r.font.bold, r.font.name = Pt(12), True, FONT_BODY
    r.font.color.rgb = RGBColor.from_string(WHITE)


def picture(slide, path, x, y, w):
    iw, ih = Image.open(path).size
    h = w * ih / iw
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return h


def main():
    summary = pd.read_csv(RES / "summary.csv")
    names = sorted(summary["text"].unique(), key=order_key)
    img_dir = RES / "slides_img"
    img_dir.mkdir(exist_ok=True)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]

    # --- title slide ---
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    text(s, 0.8, 2.2, 11.7, 1.0, "LCA creativity metric — per-text results", 40, WHITE, bold=True, font=FONT_HEAD)
    text(s, 0.8, 3.35, 11.7, 0.9,
         "WordNet LCA height of each model's top-5 next-word guesses, at every measured noun and verb. "
         "Each slide shows one text: smoothed over 10 words (top) and at every word (bottom).",
         18, "CADCFC")
    models = "  ·  ".join(summary["model"].unique())
    kinds = summary.drop_duplicates("text")["kind"]
    n_f, n_c = (kinds == "factual").sum(), (kinds == "creative").sum()
    text(s, 0.8, 4.7, 11.7, 0.4, f"Models: {models}", 14, "CADCFC")
    text(s, 0.8, 5.15, 11.7, 0.4, f"{n_f} factual texts (Wikipedia)  ·  {n_c} creative texts (Harry Potter, 1922 joke book)",
         14, "CADCFC")
    pill(s, 0.8, 6.0, 1.4, 0.38, "FACTUAL", KIND_COLOR["factual"])
    pill(s, 2.4, 6.0, 1.4, 0.38, "CREATIVE", KIND_COLOR["creative"])

    # --- one slide per text ---
    for name in names:
        rows = summary[summary["text"] == name]
        kind = rows["kind"].iloc[0]
        table, pos = load_table(name)
        img = img_dir / f"{name}.png"
        figure(name, table, img)

        s = prs.slides.add_slide(blank)
        text(s, 0.5, 0.3, 10.3, 0.55, pretty(name), 26, INK, bold=True, font=FONT_HEAD, anchor=MSO_ANCHOR.MIDDLE)
        pill(s, W - 0.5 - 1.6, 0.38, 1.6, 0.4, kind.upper(), KIND_COLOR[kind])
        avg = table.mean(axis=1)  # the black line in the plots
        stats = (f"Average of the 3 models:  mean LCA {avg.mean():.2f}   ·   variance {avg.var():.2f}"
                 f"   ·   mean |jump| {avg.diff().abs().mean():.2f}   ·   {len(table)} measured words"
                 f"   ·   {(pos == 'v').mean():.0%} verbs")
        text(s, 0.5, 0.88, 12.3, 0.3, stats, 14, MUTED)
        picture(s, img, (W - PLOT_W) / 2, 1.3, PLOT_W)

        s.notes_slide.notes_text_frame.text = "\n".join(
            f"{r.model}: mean {r.mean:.2f}, var {r.var:.2f}, mean |jump| {r.mean_abs_jump:.2f}, "
            f"noun mean {r.noun_mean:.2f}, n={r.n}" for r in rows.itertuples())

    prs.save(OUT)
    print(f"{len(prs.slides)} slides -> {OUT}")


if __name__ == "__main__":
    main()
